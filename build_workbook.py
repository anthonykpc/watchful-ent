"""
Generates HARVEST_Contact_Tracker.xlsx for the HARVEST Food Hub contact
management system. Produces a fully formatted workbook with three tabs:
  1. All Contacts
  2. Follow-Up Dashboard
  3. How to Use This Tracker
"""

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_FILE = "HARVEST_Contact_Tracker.xlsx"

# Brand colors
DEEP_GREEN = "2D6A4F"          # header fill
LIGHT_GREEN = "E8F3EC"         # alternating row fill
WHITE = "FFFFFF"
TEXT_DARK = "1B3A2B"
BORDER_GREEN = "B7D9C2"

# Column widths for the "All Contacts" sheet
ALL_CONTACTS_COLUMNS = [
    ("Timestamp", 18),
    ("First Name", 15),
    ("Last Name", 15),
    ("Organization", 28),
    ("Title/Role", 22),
    ("Email", 28),
    ("Phone", 16),
    ("Contact Category", 32),
    ("HARVEST Pillar", 22),
    ("Relationship Warmth", 20),
    ("Connection Source", 26),
    ("Last Contact Date", 18),
    ("Next Follow-Up Date", 20),
    ("Notes/Relationship History", 55),
    ("Added By", 16),
    ("Active?", 12),
]

CATEGORY_OPTIONS = [
    "Funders/Grantmakers",
    "Government/Elected Officials",
    "Community Organizations/Nonprofits",
    "Food Entrepreneurs/Potential Kitchen Members",
    "Healthcare/RWJBH Partners",
    "Peer Food Hubs",
    "Media/Press",
    "Farmers/UAC-related",
    "Vendors/Suppliers",
    "Other",
]

PILLAR_OPTIONS = [
    "Kitchen Incubator",
    "Business Accelerator",
    "Health & Wellness Hub",
    "All Three",
    "N/A",
]

WARMTH_OPTIONS = ["Hot", "Warm", "Cold", "Unknown"]
ACTIVE_OPTIONS = ["Yes", "No", "Archived"]

DASHBOARD_COLUMNS = [
    ("First Name", 15),
    ("Last Name", 15),
    ("Organization", 28),
    ("Contact Category", 32),
    ("Relationship Warmth", 20),
    ("Next Follow-Up Date", 20),
    ("Notes/Relationship History", 60),
]

MAX_DATA_ROW = 2000  # size of the pre-validated / referenced range


def header_font():
    return Font(name="Calibri", size=11, bold=True, color=WHITE)


def header_fill():
    return PatternFill("solid", fgColor=DEEP_GREEN)


def thin_border():
    side = Side(style="thin", color=BORDER_GREEN)
    return Border(left=side, right=side, top=side, bottom=side)


def build_all_contacts(wb):
    ws = wb.active
    ws.title = "All Contacts"

    # Row 1 headers
    for idx, (label, width) in enumerate(ALL_CONTACTS_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.font = header_font()
        cell.fill = header_fill()
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = thin_border()
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"

    # Auto-filter across all columns (data rows will be added later)
    last_col_letter = get_column_letter(len(ALL_CONTACTS_COLUMNS))
    ws.auto_filter.ref = f"A1:{last_col_letter}1"

    # Date formatting for Last Contact Date (L) and Next Follow-Up Date (M)
    for col_letter in ("L", "M"):
        for row in range(2, MAX_DATA_ROW + 1):
            ws[f"{col_letter}{row}"].number_format = "yyyy-mm-dd"

    # Timestamp column also formatted as date-time by default
    for row in range(2, MAX_DATA_ROW + 1):
        ws[f"A{row}"].number_format = "yyyy-mm-dd hh:mm"

    # Data validations — stored inline (all lists fit within Excel's 255-char limit)
    def add_dv(options, col_letter):
        dv = DataValidation(
            type="list",
            formula1='"{}"'.format(",".join(options)),
            allow_blank=True,
            showDropDown=False,   # False = dropdown IS shown (inverted in openpyxl)
        )
        dv.error = "Please pick a value from the dropdown."
        dv.errorTitle = "Invalid entry"
        dv.prompt = "Choose from the list."
        dv.promptTitle = "Select one"
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{MAX_DATA_ROW}")

    add_dv(CATEGORY_OPTIONS, "H")   # Contact Category
    add_dv(PILLAR_OPTIONS, "I")     # HARVEST Pillar
    add_dv(WARMTH_OPTIONS, "J")     # Relationship Warmth
    add_dv(ACTIVE_OPTIONS, "P")     # Active?

    # Alternating row shading via conditional formatting (applies as data is added)
    banded_range = f"A2:{last_col_letter}{MAX_DATA_ROW}"
    band_fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    ws.conditional_formatting.add(
        banded_range,
        FormulaRule(formula=["MOD(ROW(),2)=0"], fill=band_fill),
    )

    # Default body font + alignment for readability on the data range
    body_font = Font(name="Calibri", size=11, color=TEXT_DARK)
    body_align = Alignment(vertical="top", wrap_text=True)
    for row in range(2, MAX_DATA_ROW + 1):
        for col in range(1, len(ALL_CONTACTS_COLUMNS) + 1):
            c = ws.cell(row=row, column=col)
            c.font = body_font
            c.alignment = body_align


def build_follow_up_dashboard(wb):
    ws = wb.create_sheet("Follow-Up Dashboard")

    # Title banner in row 1 (spans all 7 dashboard columns)
    last_col_letter = get_column_letter(len(DASHBOARD_COLUMNS))
    ws.merge_cells(f"A1:{last_col_letter}1")
    title_cell = ws["A1"]
    title_cell.value = "Upcoming Follow-Ups — Next 14 Days"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    title_cell.fill = header_fill()
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Column headers in row 2
    for idx, (label, width) in enumerate(DASHBOARD_COLUMNS, start=1):
        cell = ws.cell(row=2, column=idx, value=label)
        cell.font = header_font()
        cell.fill = header_fill()
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = thin_border()
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[2].height = 30

    # Freeze the title + header rows
    ws.freeze_panes = "A3"

    # Dynamic FILTER formula (Excel 365 / Excel Online).
    # Pulls rows from All Contacts where Next Follow-Up Date (col M) is
    # between TODAY() and TODAY()+14 inclusive, skipping blanks, sorted
    # ascending by date. CHOOSE re-projects to the requested column order.
    src = f"'All Contacts'!"
    formula = (
        "=IFERROR("
        "SORT("
        "FILTER("
        "CHOOSE({1,2,3,4,5,6,7},"
        f"{src}B2:B{MAX_DATA_ROW},"
        f"{src}C2:C{MAX_DATA_ROW},"
        f"{src}D2:D{MAX_DATA_ROW},"
        f"{src}H2:H{MAX_DATA_ROW},"
        f"{src}J2:J{MAX_DATA_ROW},"
        f"{src}M2:M{MAX_DATA_ROW},"
        f"{src}N2:N{MAX_DATA_ROW}"
        "),"
        f"(ISNUMBER({src}M2:M{MAX_DATA_ROW}))"
        f"*({src}M2:M{MAX_DATA_ROW}>=TODAY())"
        f"*({src}M2:M{MAX_DATA_ROW}<=TODAY()+14)"
        "),"
        "6,1"
        "),"
        '"No follow-ups due in the next 14 days."'
        ")"
    )
    result_cell = ws.cell(row=3, column=1, value=formula)
    result_cell.font = Font(name="Calibri", size=11, color=TEXT_DARK)
    result_cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Pre-format the spill range: date column (F) + body styling for many rows
    for row in range(3, 200):
        for col in range(1, len(DASHBOARD_COLUMNS) + 1):
            c = ws.cell(row=row, column=col)
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row, column=6).number_format = "yyyy-mm-dd"

    # Banded shading on the spill area
    band_fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    ws.conditional_formatting.add(
        f"A3:{last_col_letter}200",
        FormulaRule(formula=["MOD(ROW(),2)=1"], fill=band_fill),
    )


def build_how_to_use(wb):
    ws = wb.create_sheet("How to Use This Tracker")

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110

    title = ws["B2"]
    title.value = "HARVEST Food Hub — Contact Tracker: How to Use"
    title.font = Font(name="Calibri", size=18, bold=True, color=WHITE)
    title.fill = header_fill()
    title.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B2:B2")

    subtitle = ws["B3"]
    subtitle.value = (
        "A plain-English guide for the HARVEST/UAC team. "
        "No prior CRM experience needed."
    )
    subtitle.font = Font(name="Calibri", size=11, italic=True, color=TEXT_DARK)
    subtitle.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 22

    sections = [
        (
            "1. Adding a new contact (the preferred way)",
            "Use the Microsoft Form — it's the fastest way to add a contact and "
            "guarantees the data lands in the right columns.\n\n"
            "   • Form URL:  [PASTE MICROSOFT FORM URL HERE AFTER SETUP]\n"
            "   • Open the link, fill in what you know, and submit.\n"
            "   • New submissions appear automatically as a new row on the "
            "'All Contacts' tab (usually within a minute).\n"
            "   • You don't have to fill in every field — only First Name, "
            "Last Name, and Contact Category are required. Add more detail later.",
        ),
        (
            "2. Adding a contact directly in Excel (backup method)",
            "If the form is unavailable, you can type directly into the "
            "'All Contacts' tab.\n"
            "   • Click the first empty row.\n"
            "   • Fill in First Name, Last Name, and as many other fields as "
            "you have.\n"
            "   • For the dropdown columns (Contact Category, HARVEST Pillar, "
            "Relationship Warmth, Active?), click the small arrow in the cell "
            "and pick a value — do not type free text.\n"
            "   • Put today's date in the Timestamp column (or leave blank if "
            "added via the form — the form fills it in for you).",
        ),
        (
            "3. Using AutoFilter to sort and search",
            "Each column header on 'All Contacts' has a small filter arrow.\n"
            "   • Click the arrow on 'Contact Category' to show only Funders, "
            "or only Food Entrepreneurs, etc.\n"
            "   • Click the arrow on 'HARVEST Pillar' to see only Kitchen "
            "Incubator contacts, etc.\n"
            "   • Click the arrow on 'Relationship Warmth' to see who is Hot "
            "vs. Cold.\n"
            "   • You can combine filters (e.g., Funders + Warm) to build a "
            "short call list.\n"
            "   • To clear a filter, click the arrow again and choose "
            "'Clear Filter'.",
        ),
        (
            "4. Updating a contact after a meeting or email",
            "Find the contact's row (use Ctrl+F / Cmd+F to search by name).\n"
            "   • Update 'Last Contact Date' to today.\n"
            "   • Update 'Next Follow-Up Date' to the date you plan to reach "
            "out again (leave blank if no follow-up is needed).\n"
            "   • Add a dated line to 'Notes/Relationship History' — keep "
            "older notes; don't overwrite them. Example format:\n"
            "        2026-04-22: Met at Newark City Hall; interested in "
            "vendor night. Send MOU draft.\n"
            "   • Bump 'Relationship Warmth' up or down as appropriate.",
        ),
        (
            "5. Reading the Follow-Up Dashboard tab",
            "The 'Follow-Up Dashboard' tab is read-only — it updates itself.\n"
            "   • It shows every contact whose 'Next Follow-Up Date' is "
            "between today and 14 days from now, sorted soonest first.\n"
            "   • If the list is empty, you'll see: "
            "'No follow-ups due in the next 14 days.'\n"
            "   • Check this tab at the start of each week as your working "
            "call/email list.\n"
            "   • Do NOT type or edit cells on this tab — any edits you make "
            "to data still need to happen on 'All Contacts'.",
        ),
        (
            "6. Archiving instead of deleting",
            "If a contact is no longer relevant (left their org, duplicate "
            "entry, etc.), DO NOT delete the row.\n"
            "   • Open the 'Active?' column dropdown and change it to "
            "'Archived' (or 'No').\n"
            "   • Add a note explaining why in 'Notes/Relationship History'.\n"
            "   • Why: deleting rows breaks formulas and loses history. "
            "Archiving keeps the record while hiding it from your active "
            "filters.",
        ),
        (
            "7. Sharing and permissions",
            "This file lives in SharePoint / OneDrive for Business under the "
            "RWJBarnabas Health Microsoft 365 tenant.\n"
            "   • Share it with teammates using Share > Specific People.\n"
            "   • Give Edit access to anyone who adds or updates contacts; "
            "View access for leadership who only need to read it.\n"
            "   • Do not email the .xlsx file around — always share the "
            "SharePoint/OneDrive link so everyone sees the same live data.",
        ),
        (
            "8. Who to ask for help",
            "Team lead for this tracker:  [PASTE TEAM LEAD NAME + EMAIL HERE]\n"
            "   • Contact the team lead for: new dropdown categories, broken "
            "formulas, form-to-Excel sync problems, permissions issues.",
        ),
    ]

    row = 5
    for heading, body in sections:
        h = ws.cell(row=row, column=2, value=heading)
        h.font = Font(name="Calibri", size=13, bold=True, color=DEEP_GREEN)
        h.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22
        row += 1

        b = ws.cell(row=row, column=2, value=body)
        b.font = Font(name="Calibri", size=11, color=TEXT_DARK)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        # Approximate row height based on line count
        line_count = body.count("\n") + 1
        ws.row_dimensions[row].height = max(18 * line_count, 40)
        row += 2  # blank spacer row between sections

    # Footer reminder
    footer = ws.cell(
        row=row + 1,
        column=2,
        value=(
            "Remember: the tracker is only as useful as the notes you put in it. "
            "A one-line note after every meeting keeps HARVEST's relationships "
            "alive across staff turnover."
        ),
    )
    footer.font = Font(name="Calibri", size=11, italic=True, color=DEEP_GREEN)
    footer.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row + 1].height = 40


def main():
    wb = Workbook()
    build_all_contacts(wb)
    build_follow_up_dashboard(wb)
    build_how_to_use(wb)

    # Make "All Contacts" the default visible sheet on open
    wb.active = 0
    wb.save(OUTPUT_FILE)
    print(f"Wrote {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
